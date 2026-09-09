# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) báo cáo BÀI TẬP 1 
với công thức toán được render đẹp bằng OMML (Office Math Markup Language).
"""

import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from lxml import etree
import latex2mathml.converter

# ============================================================
# HELPER: Chuyển LaTeX → MathML → OMML → chèn vào paragraph
# ============================================================

# XSLT để chuyển MathML sang OMML (rút gọn + tự viết)
MML2OMML_XSL = r'''<?xml version="1.0" encoding="UTF-8"?>
<xsl:stylesheet version="1.0"
    xmlns:xsl="http://www.w3.org/1999/XSL/Transform"
    xmlns:mml="http://www.w3.org/1998/Math/MathML"
    xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">

  <xsl:output method="xml" encoding="UTF-8" indent="yes"/>

  <xsl:template match="mml:math">
    <m:oMath>
      <xsl:apply-templates/>
    </m:oMath>
  </xsl:template>

  <xsl:template match="mml:mrow">
    <xsl:apply-templates/>
  </xsl:template>

  <xsl:template match="mml:mi">
    <m:r>
      <m:rPr><m:sty m:val="i"/></m:rPr>
      <m:t><xsl:value-of select="."/></m:t>
    </m:r>
  </xsl:template>

  <xsl:template match="mml:mn">
    <m:r>
      <m:t><xsl:value-of select="."/></m:t>
    </m:r>
  </xsl:template>

  <xsl:template match="mml:mo">
    <m:r>
      <m:t><xsl:value-of select="."/></m:t>
    </m:r>
  </xsl:template>

  <xsl:template match="mml:mspace">
    <m:r>
      <m:t xml:space="preserve"> </m:t>
    </m:r>
  </xsl:template>

  <xsl:template match="mml:msup">
    <m:sSup>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sup><xsl:apply-templates select="*[2]"/></m:sup>
    </m:sSup>
  </xsl:template>

  <xsl:template match="mml:msub">
    <m:sSub>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sub><xsl:apply-templates select="*[2]"/></m:sub>
    </m:sSub>
  </xsl:template>

  <xsl:template match="mml:mfrac">
    <m:f>
      <m:num><xsl:apply-templates select="*[1]"/></m:num>
      <m:den><xsl:apply-templates select="*[2]"/></m:den>
    </m:f>
  </xsl:template>

  <xsl:template match="mml:msqrt">
    <m:rad>
      <m:radPr><m:degHide m:val="1"/></m:radPr>
      <m:deg/>
      <m:e><xsl:apply-templates/></m:e>
    </m:rad>
  </xsl:template>

  <xsl:template match="mml:msubsup">
    <m:sSubSup>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sub><xsl:apply-templates select="*[2]"/></m:sub>
      <m:sup><xsl:apply-templates select="*[3]"/></m:sup>
    </m:sSubSup>
  </xsl:template>

  <xsl:template match="mml:mover">
    <m:acc>
      <m:accPr><m:chr m:val="&#x0302;"/></m:accPr>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
    </m:acc>
  </xsl:template>

  <xsl:template match="mml:mtext">
    <m:r>
      <m:rPr><m:sty m:val="p"/></m:rPr>
      <m:t><xsl:value-of select="."/></m:t>
    </m:r>
  </xsl:template>

  <xsl:template match="mml:mtable">
    <m:m>
      <xsl:for-each select="mml:mtr">
        <m:mr>
          <xsl:for-each select="mml:mtd">
            <m:e><xsl:apply-templates/></m:e>
          </xsl:for-each>
        </m:mr>
      </xsl:for-each>
    </m:m>
  </xsl:template>

  <!-- Fallback: copy text -->
  <xsl:template match="text()">
    <xsl:if test="normalize-space(.) != ''">
      <m:r>
        <m:t><xsl:value-of select="."/></m:t>
      </m:r>
    </xsl:if>
  </xsl:template>

</xsl:stylesheet>'''

_xsl_tree = etree.fromstring(MML2OMML_XSL.encode('utf-8'))
_xsl_transform = etree.XSLT(_xsl_tree)


def latex_to_omml(latex_str):
    """Chuyển chuỗi LaTeX thành element OMML để chèn vào docx."""
    try:
        mathml_str = latex2mathml.converter.convert(latex_str)
        mathml_tree = etree.fromstring(mathml_str.encode('utf-8'))
        omml_tree = _xsl_transform(mathml_tree)
        return omml_tree.getroot()
    except Exception as e:
        # Fallback: trả về plain text run
        print(f"[WARN] Không chuyển được LaTeX '{latex_str}': {e}")
        return None


def add_math_paragraph(doc, latex_str, alignment=WD_ALIGN_PARAGRAPH.CENTER):
    """Thêm 1 paragraph chứa công thức toán (từ LaTeX) vào document."""
    omml = latex_to_omml(latex_str)
    para = doc.add_paragraph()
    para.alignment = alignment
    if omml is not None:
        # Wrap trong oMathPara để căn giữa
        omml_para = OxmlElement('m:oMathPara')
        omml_para.append(omml)
        para._element.append(omml_para)
    else:
        # Fallback: in plain text
        para.add_run(latex_str)
    return para


def add_inline_math(paragraph, latex_str):
    """Chèn công thức toán inline vào paragraph hiện có."""
    omml = latex_to_omml(latex_str)
    if omml is not None:
        paragraph._element.append(omml)
    else:
        run = paragraph.add_run(latex_str)
        run.font.name = 'Cambria Math'


def set_cell_text(cell, text, bold=False, alignment=WD_ALIGN_PARAGRAPH.LEFT, font_size=10):
    """Set text cho cell trong bảng."""
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = alignment
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    return para


def set_cell_math(cell, latex_str, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """Set công thức toán cho cell trong bảng."""
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = alignment
    add_inline_math(para, latex_str)
    return para


def add_styled_paragraph(doc, text, style='Normal', bold=False, font_size=11, space_after=6):
    """Thêm paragraph với style tùy chỉnh."""
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    para.paragraph_format.space_after = Pt(space_after)
    return para


def shade_cells(row, color="D9E2F3"):
    """Tô màu nền cho hàng trong bảng."""
    for cell in row.cells:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), color)
        shading.set(qn('w:val'), 'clear')
        cell._element.get_or_add_tcPr().append(shading)


# ============================================================
# TẠO TÀI LIỆU WORD
# ============================================================

def create_report():
    doc = Document()
    
    # ---- Style mặc định ----
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)
    style.paragraph_format.line_spacing = 1.5
    
    # ====== TRANG BÌA ======
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    run = title.add_run("BÁO CÁO KẾT QUẢ BÀI TẬP 1")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Môn học: Cơ sở Logic và Toán rời rạc / Đại số máy tính")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)

    doc.add_paragraph()  # Khoảng trắng

    # Bảng danh sách nhóm
    tbl = doc.add_table(rows=4, cols=4)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["STT", "Họ và tên", "Mã Học Viên / MSSV", "Lớp"]
    for i, h in enumerate(headers):
        set_cell_text(tbl.rows[0].cells[i], h, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cells(tbl.rows[0], "1F497D")
    for cell in tbl.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    data = [
        ["1", "Nguyễn Đình Duy", "(Điền MSSV)", "(Điền Lớp)"],
        ["2", "Thành viên 2", "(Điền MSSV)", "(Điền Lớp)"],
        ["3", "Thành viên 3", "(Điền MSSV)", "(Điền Lớp)"],
    ]
    for ri, row_data in enumerate(data):
        for ci, val in enumerate(row_data):
            set_cell_text(tbl.rows[ri+1].cells[ci], val, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        if ri % 2 == 0:
            shade_cells(tbl.rows[ri+1], "D6E4F0")
    
    doc.add_page_break()

    # ====== BÀI 1 ======
    doc.add_heading("BÀI 1: Biểu diễn logic vị từ", level=1)
    
    doc.add_heading("Đề bài", level=2)
    p = doc.add_paragraph("Đặt ")
    add_inline_math(p, r"C(x)")
    p.add_run(': "x có một con mèo", ')
    add_inline_math(p, r"D(x)")
    p.add_run(': "x có một con chó", ')
    add_inline_math(p, r"F(x)")
    p.add_run(': "x có một con chồn".')
    
    doc.add_paragraph("Biểu diễn các phát biểu sau theo C(x), D(x), F(x), các lượng từ và các phép nối logic. Xét không gian biến là các sinh viên trong lớp.")
    
    doc.add_heading("Lời giải chi tiết", level=2)
    
    bai1_items = [
        {
            "cau": "a) Một sinh viên trong lớp có một con mèo, một con chó hay một con chồn.",
            "phan_tich": "Phát biểu khẳng định sự tồn tại (∃) của ít nhất một sinh viên sở hữu ít nhất một trong ba con vật (phép tuyển ∨).",
            "latex": r"\exists x \, (C(x) \lor D(x) \lor F(x))"
        },
        {
            "cau": "b) Tất cả sinh viên trong lớp có một con mèo, một con chó hay một con chồn.",
            "phan_tich": "Khẳng định với mọi (∀) sinh viên trong lớp, mỗi người đều có ít nhất một con vật nuôi.",
            "latex": r"\forall x \, (C(x) \lor D(x) \lor F(x))"
        },
        {
            "cau": "c) Một sinh viên nào đó có một con mèo và một con chồn nhưng không có chó.",
            "phan_tich": "Tồn tại sinh viên x thỏa mãn đồng thời: C(x) đúng, F(x) đúng và phủ định ¬D(x).",
            "latex": r"\exists x \, (C(x) \land F(x) \land \neg D(x))"
        },
        {
            "cau": "d) Không có sinh viên nào trong lớp có một con mèo, một con chó và một con chồn.",
            "phan_tich": "Phủ định mệnh đề tồn tại. Tương đương (De Morgan): ∀x ¬(C(x) ∧ D(x) ∧ F(x)).",
            "latex": r"\neg \exists x \, (C(x) \land D(x) \land F(x))"
        },
        {
            "cau": "e) Với mỗi loại con vật trên, có một sinh viên trong lớp có một con.",
            "phan_tich": "Với từng loài, tồn tại ít nhất một sinh viên nuôi con vật đó (không bắt buộc cùng một sinh viên).",
            "latex": r"(\exists x \, C(x)) \land (\exists y \, D(y)) \land (\exists z \, F(z))"
        }
    ]
    
    for item in bai1_items:
        doc.add_heading(item["cau"], level=3)
        doc.add_paragraph(f"Phân tích: {item['phan_tich']}")
        p = doc.add_paragraph("Công thức logic:")
        p.runs[0].bold = True
        add_math_paragraph(doc, item["latex"])
    
    doc.add_page_break()

    # ====== BÀI 2* ======
    doc.add_heading("BÀI 2*: Logic bậc nhất nâng cao và Chứng minh hình thức", level=1)
    
    doc.add_heading("Đề bài", level=2)
    p = doc.add_paragraph("Đặt: L(x): \"x là nhà logic\", C(x): \"x uống café\", W(x): \"x làm việc chăm chỉ\", T(x): \"x phát biểu định lý\", f(x): hàm trả ra người bạn duy nhất của x.")

    doc.add_heading("Phần 1: Biểu diễn logic bậc nhất (có sử dụng dấu =)", level=2)
    
    bai2_items = [
        {
            "cau": "a. Không nhà logic nào uống café",
            "note": "Dạng chuẩn: ∀x (L(x) → ¬C(x)).",
            "latex": r"\forall x \forall y \, ((y = x \land L(x)) \to \neg C(y))"
        },
        {
            "cau": "b. Bất kỳ ai là một nhà logic cũng đều là bạn của ai đó",
            "note": "Nếu x là nhà logic thì tồn tại y sao cho bạn của y là x.",
            "latex": r"\forall x \, (L(x) \to \exists y \, (x = f(y)))"
        },
        {
            "cau": "c. Không ai phát biểu được định lý lại có một người bạn uống café",
            "note": "Dạng rút gọn: ∀x (T(x) → ¬C(f(x))).",
            "latex": r"\forall x \forall y \, ((y = f(x) \land T(x)) \to \neg C(y))"
        },
        {
            "cau": "d. Ai có bạn làm việc chăm chỉ thì hoặc là nhà logic hoặc cũng chăm chỉ",
            "note": "Tương đương: ∀x (W(f(x)) → (L(x) ∨ W(x))).",
            "latex": r"\forall x \forall y \, ((y = f(x) \land W(y)) \to (L(x) \lor W(x)))"
        },
        {
            "cau": "e. Mọi người bạn là một nhà logic",
            "note": "Tương đương: ∀y L(f(y)).",
            "latex": r"\forall x \, (\exists y \, (x = f(y)) \to L(x))"
        }
    ]
    
    for item in bai2_items:
        doc.add_heading(item["cau"], level=3)
        doc.add_paragraph(item["note"])
        p = doc.add_paragraph("Biểu diễn (sử dụng dấu =):")
        p.runs[0].bold = True
        add_math_paragraph(doc, item["latex"])
    
    # Phần 2: Chứng minh
    doc.add_heading("Phần 2: Chứng minh hình thức suy luận logic", level=2)
    
    doc.add_paragraph("Các tiền đề:").runs[0].bold = True
    p1 = doc.add_paragraph("(1) Tất cả nhà logic đều uống café:  ")
    add_inline_math(p1, r"\forall x \, (L(x) \to C(x))")
    p2 = doc.add_paragraph("(2) Ai không phát biểu được định lý thì không uống café:  ")
    add_inline_math(p2, r"\forall x \, (\neg T(x) \to \neg C(x))")
    p3 = doc.add_paragraph("(3) Có một số người mà bạn của họ là nhà logic:  ")
    add_inline_math(p3, r"\exists x \, L(f(x))")
    
    p_kl = doc.add_paragraph("Kết luận cần chứng minh: ")
    p_kl.runs[0].bold = True
    add_inline_math(p_kl, r"\exists x \, (L(x) \land C(x) \land T(x))")

    # Bảng bước chứng minh
    doc.add_paragraph()
    proof_steps = [
        ("1", "∃x L(f(x))", "Tiền đề (3)"),
        ("2", "L(f(c))", "Đặc tả tồn tại (EI) từ 1, c là cá thể cụ thể"),
        ("3", "Đặt y₀ = f(c). Ta có L(y₀)", "f(c) là phần tử xác định thuộc miền biến"),
        ("4", "∀x (L(x) → C(x))", "Tiền đề (1)"),
        ("5", "L(y₀) → C(y₀)", "Đặc tả phổ dụng (UI) từ 4 thay x = y₀"),
        ("6", "C(y₀)", "Modus Ponens từ 3 và 5"),
        ("7", "∀x (¬T(x) → ¬C(x))", "Tiền đề (2)"),
        ("8", "∀x (C(x) → T(x))", "Phản đảo (Contrapositive) tương đương từ 7"),
        ("9", "C(y₀) → T(y₀)", "Đặc tả phổ dụng (UI) từ 8 thay x = y₀"),
        ("10", "T(y₀)", "Modus Ponens từ 6 và 9"),
        ("11", "L(y₀) ∧ C(y₀) ∧ T(y₀)", "Luật hội (Conjunction) từ 3, 6 và 10"),
        ("12", "∃x (L(x) ∧ C(x) ∧ T(x))", "Khái quát hóa tồn tại (EG) từ 11"),
    ]
    
    proof_tbl = doc.add_table(rows=len(proof_steps)+1, cols=3)
    proof_tbl.style = 'Table Grid'
    proof_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for i, h in enumerate(["Bước", "Biểu thức khẳng định", "Lý do / Quy tắc suy diễn"]):
        set_cell_text(proof_tbl.rows[0].cells[i], h, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
    shade_cells(proof_tbl.rows[0], "1F497D")
    for cell in proof_tbl.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    for ri, (step, expr, reason) in enumerate(proof_steps):
        set_cell_text(proof_tbl.rows[ri+1].cells[0], step, alignment=WD_ALIGN_PARAGRAPH.CENTER, font_size=10)
        set_cell_text(proof_tbl.rows[ri+1].cells[1], expr, font_size=10)
        set_cell_text(proof_tbl.rows[ri+1].cells[2], reason, font_size=10)
        if ri % 2 == 0:
            shade_cells(proof_tbl.rows[ri+1], "D6E4F0")
    
    p_qed = doc.add_paragraph()
    p_qed.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p_qed.add_run("⟹ Điều phải chứng minh (Q.E.D.) ■")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_page_break()

    # ====== BÀI 3 ======
    doc.add_heading("BÀI 3: Giải và biện luận phương trình theo tham số m", level=1)
    
    # 3a
    doc.add_heading("a. Phương trình bậc nhất theo tham số m", level=2)
    p = doc.add_paragraph("Ví dụ đề bài: Giải và biện luận:  ")
    p.runs[0].bold = True
    add_inline_math(p, r"(m^2 - 4)x + 2m - 4 = 0")
    
    doc.add_heading("Phương pháp giải tổng quát", level=3)
    p = doc.add_paragraph("Phương trình có dạng  ")
    add_inline_math(p, r"A(m) \cdot x + B(m) = 0")
    p.add_run("  ⟺  ")
    add_inline_math(p, r"A(m) \cdot x = -B(m)")
    
    doc.add_paragraph("• Tìm các giá trị m làm A(m) = 0.")
    doc.add_paragraph("• Nếu −B(m₀) = 0: Phương trình 0x = 0 ⟹ Vô số nghiệm ∈ ℝ.")
    doc.add_paragraph("• Nếu −B(m₀) ≠ 0: Phương trình 0x = c ≠ 0 ⟹ Vô nghiệm.")
    p = doc.add_paragraph("• Nếu A(m) ≠ 0: Nghiệm duy nhất  ")
    add_inline_math(p, r"x = -\frac{B(m)}{A(m)}")
    
    doc.add_heading("Lời giải cụ thể", level=3)
    
    p = doc.add_paragraph("Hệ số  ")
    add_inline_math(p, r"A(m) = m^2 - 4 = 0")
    p.add_run("  ⟺  m = 2 hoặc m = −2.")
    
    p = doc.add_paragraph("")
    run = p.add_run("Nếu m = 2: ")
    run.bold = True
    p.add_run("Phương trình: 0·x + 0 = 0 ⟹ ")
    run2 = p.add_run("Vô số nghiệm thuộc ℝ.")
    run2.bold = True
    run2.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    
    p = doc.add_paragraph("")
    run = p.add_run("Nếu m = −2: ")
    run.bold = True
    p.add_run("Phương trình: 0·x + (−8) = 0 ⟹ 0x = 8 ⟹ ")
    run2 = p.add_run("Vô nghiệm.")
    run2.bold = True
    run2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    
    p = doc.add_paragraph("")
    run = p.add_run("Nếu m ≠ 2 và m ≠ −2: ")
    run.bold = True
    p.add_run("Phương trình có nghiệm duy nhất:")
    add_math_paragraph(doc, r"x = \frac{-(2m - 4)}{m^2 - 4} = \frac{-2(m - 2)}{(m - 2)(m + 2)} = \frac{-2}{m + 2}")
    
    # 3b*
    doc.add_heading("b*. Phương trình bậc hai theo tham số m", level=2)
    p = doc.add_paragraph("Ví dụ đề bài: Giải và biện luận:  ")
    p.runs[0].bold = True
    add_inline_math(p, r"(m - 2)x^2 - 2(m + 2)x + m + 2 = 0")
    
    # TH1
    p = doc.add_paragraph("")
    run = p.add_run("Trường hợp 1: A = 0 ⟺ m − 2 = 0 ⟺ m = 2.")
    run.bold = True
    doc.add_paragraph("Phương trình trở thành bậc nhất: −8x + 4 = 0 ⟹ x = 1/2.")

    # TH2
    p = doc.add_paragraph("")
    run = p.add_run("Trường hợp 2: A ≠ 0 ⟺ m ≠ 2.")
    run.bold = True
    p = doc.add_paragraph("Tính biệt thức thu gọn Δ':")
    add_math_paragraph(doc, r"\Delta' = (B')^2 - AC = [-(m+2)]^2 - (m-2)(m+2) = 4(m+2)")
    
    p = doc.add_paragraph("• ")
    run = p.add_run("Nếu Δ' < 0 ⟺ m < −2: ")
    run.bold = True
    run2 = p.add_run("Phương trình vô nghiệm thực.")
    run2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    
    p = doc.add_paragraph("• ")
    run = p.add_run("Nếu Δ' = 0 ⟺ m = −2: ")
    run.bold = True
    p.add_run("Nghiệm kép x = 0.")
    
    p = doc.add_paragraph("• ")
    run = p.add_run("Nếu Δ' > 0 ⟺ m > −2 (và m ≠ 2): ")
    run.bold = True
    p.add_run("Hai nghiệm phân biệt:")
    add_math_paragraph(doc, r"x_{1,2} = \frac{(m + 2) \pm 2\sqrt{m + 2}}{m - 2}")

    doc.add_page_break()

    # ====== BÀI 4 ======
    doc.add_heading("BÀI 4: Phân tích đa thức thành nhân tử từng bước", level=1)
    
    doc.add_heading("Yêu cầu và Thuật toán", level=2)
    doc.add_paragraph("Viết chương trình thực hiện giải từng bước bài toán phân tích đa thức thành nhân tử:")
    doc.add_paragraph("• Sử dụng hàm factor() của thư viện SymPy để thu được kết quả cuối cùng.")
    doc.add_paragraph("• Suy luận ngược từ kết quả: Tách các hạng tử của thừa số thứ nhất để nhân với thừa số còn lại.")
    doc.add_paragraph("• Khai triển thành các nhóm trung gian, sau đó đặt nhân tử chung từng nhóm và gom lại thành tích.")
    doc.add_paragraph("• Sử dụng đệ quy cho trường hợp có từ 3 thừa số trở lên.")
    
    doc.add_heading("Ví dụ mẫu", level=2)
    p = doc.add_paragraph("Phân tích đa thức  ")
    add_inline_math(p, r"x^2 - y^2")
    p.add_run(":")
    
    add_math_paragraph(doc, r"x^2 - y^2 = (x^2 - xy) + (xy - y^2)")
    add_math_paragraph(doc, r"= x(x - y) + y(x - y)")
    add_math_paragraph(doc, r"= (x + y)(x - y)")
    
    doc.add_heading("Mã nguồn Python (trích đoạn)", level=2)
    
    code_text = '''import sympy as sp

def phan_tich_tung_buoc_2_thua_so(f1, f2):
    terms_f1 = sp.Add.make_args(f1)
    
    # Bước 1: Khai triển từng nhóm
    expanded_groups = [sp.expand(t * f2) for t in terms_f1]
    step1_str = " + ".join(f"({g})" for g in expanded_groups)
    
    # Bước 2: Đặt nhân tử chung cho từng nhóm
    step2_str = " + ".join(f"{t}({f2})" for t in terms_f1)
    
    # Bước 3: Nhân tử tích cuối cùng
    step3_str = f"({f1})({f2})"
    
    return step1_str, step2_str, step3_str'''
    
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    # Tô nền xám nhạt
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F2F2F2')
    shading.set(qn('w:val'), 'clear')
    p._element.get_or_add_pPr().append(shading)

    doc.add_page_break()

    # ====== HƯỚNG DẪN CHẠY ======
    doc.add_heading("Hướng dẫn chạy mã nguồn và Cấu trúc nộp bài", level=1)
    
    doc.add_heading("Cấu trúc thư mục nộp bài", level=2)
    structure = """CDCSTT/
├── Bao_Cao_Bai_Tap_1.tex          # Báo cáo LaTeX
├── Bao_Cao_Bai_Tap_1.docx         # Báo cáo Word
├── Danh_sach_nhom.csv              # Danh sách nhóm
├── bai_1_logic.py                  # Mã nguồn bài 1
├── bai_2_logic_chung_minh.py       # Mã nguồn bài 2*
├── bai_3_bien_luan_pt.py           # Mã nguồn bài 3
├── bai_4_phan_tich_nhan_tu.py      # Mã nguồn bài 4
├── main.py                         # Menu chạy tổng hợp
├── requirements.txt                # Thư viện phụ thuộc
└── README.md                       # Tài liệu hướng dẫn"""
    p = doc.add_paragraph()
    run = p.add_run(structure)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F2F2F2')
    shading.set(qn('w:val'), 'clear')
    p._element.get_or_add_pPr().append(shading)
    
    doc.add_heading("Cách cài đặt và thực thi", level=2)
    doc.add_paragraph("1. Cài đặt thư viện: pip install -r requirements.txt")
    doc.add_paragraph("2. Chạy toàn bộ bài tập: python main.py --all")
    doc.add_paragraph("3. Chạy giao diện tương tác: python main.py")
    doc.add_paragraph("4. Chạy riêng lẻ từng bài:")
    doc.add_paragraph("    • python bai_1_logic.py")
    doc.add_paragraph("    • python bai_2_logic_chung_minh.py")
    doc.add_paragraph("    • python bai_3_bien_luan_pt.py")
    doc.add_paragraph("    • python bai_4_phan_tich_nhan_tu.py")
    
    # Lưu file
    output_path = "Bao_Cao_Bai_Tap_1.docx"
    doc.save(output_path)
    print(f"[OK] Đã tạo file: {output_path}")
    return output_path


if __name__ == "__main__":
    create_report()
